# courseadmin.py — the plumbing under test-admin: paths, the report, the readers
# Usage: imported by test-admin; not run directly.
#
# Not a program; there is nothing to run here. test-admin produces every piece
# of paper a test needs -- the per-section attendance sheets, the Print Services
# order, and the envelope split -- and this is the half of it that is not about
# any one of those: where things are, how the run reports on itself, and how
# each data file is read.
#
# WHY IT IS A SEPARATE FILE. It began as what two scripts had in common,
# make-attendance-sheets and make-print-job. They are one script now, and this
# stayed apart because the division still earns its keep: test-admin reads as
# the job, top to bottom, and this reads as the facts the job stands on.
#
# WHY IT CAN BE A PLAIN SIBLING IMPORT. Python puts a script's own directory on
# sys.path[0] before anything else -- so `import courseadmin' resolves here with
# no path manipulation, no package, and no installation.
#
# WHAT BELONGS HERE is what does not depend on the paperwork being made.
# Deliberately not here: latex_escape and check_typesettable (about typesetting
# student names), compile_sheet, and every piece of the booklet arithmetic.
#
# PORTING TO ANOTHER COURSE: nothing here is edited. This file is shared across
# courses (it lives in the machinery submodule), so the two facts that differ
# from one offering to the next -- the course number, and the Team's OneDrive
# folder that the class lists, sac-students.xlsx and the output hang off -- are
# read from the course's own course-info/admin-settings.csv. Everything else
# adapts on its own from course-sections.csv.

import argparse
import csv
import os
import sys
import unicodedata
from collections import OrderedDict
from pathlib import Path

# THE ONE THIRD-PARTY IMPORT, and it is guarded on purpose. Everything else here
# is stdlib; openpyxl arrives through the uv inline block at the top of each
# script, so importing this module in a bare `python3' for a quick poke would
# otherwise die on a traceback that says nothing useful. Guarded, the failure is
# deferred to the one function that needs it -- see read_sac_students.
try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

# WHERE THINGS ARE. The starting point is this file's own location and never the
# cwd, so the scripts work however they are invoked. But its location is not
# itself the answer: the generator has lived in _admin/, in a subdirectory of it,
# in .course-machinery-local/, and now in the shared machinery at
# .course-machinery/admin/. So the COURSE is found by walking up from here, and
# everything else is named relative to that.
SCRIPT_DIR = Path(__file__).resolve().parent

# The course's data folder, by name. It holds the facts about this offering --
# sections, rooms, dates, the print order -- and is visible because colleagues
# edit it. The same folder is on TEXINPUTS, which is how the LaTeX packages
# find the very same files (see .course-machinery/latexmkrc-shared).
DATA_DIR_NAME = "course-info"

# The name to put in front of an error message. Taken from argv rather than
# hardcoded, so that the program which was run owns its errors:
# `test-admin: error: ...', not this file's name.
PROGRAM = Path(sys.argv[0]).name or "courseadmin"


def find_repo_root():
    """The course repository this script was run from -> a directory, or None.

    The course is the first ancestor that holds a course-info/ folder.

    NOT "THE FIRST ANCESTOR HOLDING .git", which is what this used to ask. This
    file now lives in a git SUBMODULE (.course-machinery/), and a submodule has
    a .git of its own -- so that test stopped one level too early and named the
    machinery as the course. Nothing would have failed loudly: the data folder
    would simply not have been found. A course-info/ folder is the thing a
    course has and the machinery does not, so it is the honest landmark.

    The walk still stops at the first hit, so a script run from one course's
    checkout of the machinery can only ever find that course.
    """
    for directory in [SCRIPT_DIR, *SCRIPT_DIR.parents]:
        if (directory / DATA_DIR_NAME).is_dir():
            return directory
    return None


def find_data_dir(repo_root):
    """The course's data folder, <repo>/course-info -> a directory, or None.

    NAMED, NOT SEARCHED FOR. This used to walk up from the script looking for
    course-sections.csv, which worked while the data sat beside the scripts or
    above them. It is now a SIBLING of the scripts' directory, and a walk
    upward never looks sideways, so the folder is named outright -- the same
    way BUILD_DIR is, below, and for the same reason.

    ONLY EVER INSIDE THIS REPOSITORY. Looking further afield could pick up
    ANOTHER course's course-sections.csv, and the sheets would then be built
    for the wrong course's sections without a word of complaint. Better to find
    nothing and say so, which is what returning None arranges.

    course-sections.csv is the test for "this is the right folder" because it
    is the course's definition: it says which sections exist, in what order,
    and who teaches them.
    """
    if repo_root is None:
        return None
    candidate = repo_root / DATA_DIR_NAME
    if (candidate / "course-sections.csv").is_file():
        return candidate
    return None


REPO_ROOT = find_repo_root()
DATA_DIR = find_data_dir(REPO_ROOT)
SECTIONS_CSV = (DATA_DIR or SCRIPT_DIR) / "course-sections.csv"

# Where the attendance sheet is BUILT: TESTS/ADMIN/ in the course tree, the same
# place the Team keeps it (Tests/ADMIN), beside the tests themselves. It
# is a scratch directory now -- the generated .tex is written and compiled here,
# and the PDF is then moved to OUTPUT_DIR (below), so between runs it holds only
# the .latexmkrc a build needs. Built in the repo rather than in OUTPUT_DIR
# because .latexmkrc finds the machinery by walking up to .course-machinery/,
# which it can only do from inside the course tree.
#
# NAMED OUTRIGHT, not derived from the data directory. It used to be
# `DATA_DIR / "test-docs"', which was true only while the data sat in _admin/
# beside its own output subdirectory. When the data moved that expression
# quietly followed it and pointed at a directory that did not exist -- the
# generator would have created a test-docs/ beside the data and written student
# names into it. The output directory and the data directory are separate facts
# now, so this states it rather than inferring it.
#
# The attendance sheet's layout is attendance.sty, which no longer has to live in
# this directory: it sits in .course-machinery/admin/ and is found by name on
# TEXINPUTS, like everything else the documents load.
BUILD_DIR = (REPO_ROOT / "TESTS" / "ADMIN") if REPO_ROOT else SCRIPT_DIR


def find_mirror_dir():
    """Where build-pdfs sends built PDFs -> a directory, or None.

    The repo's top-level .MIRRORDIR names it. That file is git-ignored and
    differs per machine, which is the point of reading it: the mirror is in
    iCloud on one computer and could be anywhere on a colleague's, and a path
    written into this script would be right for exactly one person.

    THE FILE HAS OTHER LINES IN IT, and this reads it the way build-pdfs'
    parse_dest_file does so the two agree on which one is the mirror:

      blank lines, # comments      ignored
      BOTH / COPY                  a keyword about keeping a local copy
      `Label: /some/path'          a named destination for `%! copies:'
      anything else                THE MIRROR -- the first such line wins

    The mirror repeats the repo's own tree, so the tests are in its TESTS/.
    None when there is no .MIRRORDIR, it names no mirror, or the directory it
    names is not there (iCloud or OneDrive not signed in, most likely).
    """
    if REPO_ROOT is None:
        return None
    mirrordir_file = REPO_ROOT / ".MIRRORDIR"
    if not mirrordir_file.is_file():
        return None

    for line in mirrordir_file.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        if line.casefold() in ("both", "copy"):
            continue
        # A label line is a word, a colon, then a space: `Tests: /path'. The
        # space is what tells it from a path that merely contains a colon.
        label, colon, rest = line.partition(":")
        is_label = (
            colon
            and rest[:1].isspace()
            and label.replace("-", "").replace("_", "").isalnum()
        )
        if is_label:
            continue
        mirror = Path(line).expanduser()
        return mirror if mirror.is_dir() else None
    return None


# ---------------------------------------------------------------------------
# PORTING: the team's shared OneDrive folder, named for the course and the term.
#
# Built from Path.home() rather than a literal path so the username is not baked
# into a tracked file; the rest is the team's OneDrive layout and is the same on
# both machines. Three things in it are what the scripts actually touch:
#
#   General/exported-classlists/   the registrar's exports
#   General/sac-students.xlsx      who writes at the SAC
#   Tests/ADMIN/                   everything test-admin produces
#
# NEITHER FALLS BACK TO A LOCAL COPY, and for the same reason: a stale local file
# quietly standing in for the shared one is the exact failure a shared folder
# exists to avoid. A missing class list stops the sheet for that instructor's
# sections; a missing sac-students.xlsx falls back to the ballpark, which
# over-prints -- see read_sac_students.
#
# WHY A WORKBOOK AND NOT A CSV, since everything else here is CSV: this is the
# one shared file the OTHER THREE INSTRUCTORS edit themselves. A .csv in a Teams
# document library opens read-only in Excel and offers only to convert itself,
# so it cannot be edited in place; an .xlsx opens for editing in the browser and
# co-authors, so four people can add rows without producing the sync-conflict
# duplicates a shared .csv would. The cost is this file's one dependency.
# ---------------------------------------------------------------------------

# THE FACTS THAT DIFFER FROM ONE OFFERING TO THE NEXT live in the course, not
# here: course-info/admin-settings.csv, two columns, `setting' and `value'.
#
#   course_number   1003 -- used in the paperwork's headings and to recognise
#                   the registrar's exports, whose file names carry it
#   team_dir        the Team's OneDrive folder, written relative to your home
#                   directory; it is named for the course and the term
#
# They were constants in these scripts while the scripts belonged to one
# course. Now that the scripts are shared, a constant here would be wrong for
# every course but one.
SETTINGS_CSV = (DATA_DIR or SCRIPT_DIR) / "admin-settings.csv"
REQUIRED_SETTINGS = ("course_number", "team_dir")


def read_admin_settings():
    """admin-settings.csv -> {setting: value}. Stops the program if it cannot.

    Runs once, when this module is imported, because the paths below are built
    from what it returns. That is too early to use die() (defined further
    down), so a failure here prints its own message and exits the same way.
    """
    def stop(message):
        print(f"{PROGRAM}: error: {message}", file=sys.stderr)
        sys.exit(1)

    if DATA_DIR is None:
        stop(
            f"no {DATA_DIR_NAME}/ folder holding course-sections.csv was found "
            f"above {SCRIPT_DIR}, so there is no course to work on. Run this "
            f"from a course's own checkout of the machinery."
        )
    if not SETTINGS_CSV.is_file():
        stop(
            f"{SETTINGS_CSV} not found. It needs a header line `setting,value' "
            f"and one row for each of: {', '.join(REQUIRED_SETTINGS)}."
        )

    settings = {}
    with SETTINGS_CSV.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            name = (row.get("setting") or "").strip()
            if name:
                settings[name] = (row.get("value") or "").strip()

    missing = [name for name in REQUIRED_SETTINGS if not settings.get(name)]
    if missing:
        stop(
            f"{SETTINGS_CSV.name} has no value for: {', '.join(missing)}. "
            f"It needs a header line `setting,value' and a row for each."
        )
    return settings


SETTINGS = read_admin_settings()
COURSE_NUMBER = SETTINGS["course_number"]
TEAM_DIR = Path.home() / SETTINGS["team_dir"]
SHARED_DIR = TEAM_DIR / "General"
DEFAULT_ROSTER_DIR = SHARED_DIR / "exported-classlists"
SAC_STUDENTS_XLSX = SHARED_DIR / "sac-students.xlsx"

# Where the finished paperwork goes: the Team's Tests/ADMIN, beside the tests
# themselves. The attendance-sheet PDF and the test's workbook land here
# and nowhere else -- not in the repo, not in the iCloud mirror. The Team is
# shared only with the other instructors, who see all of this anyway.
OUTPUT_DIR = TEAM_DIR / "Tests" / "ADMIN"

# The file the workbook replaced. Not read -- only noticed, so that a leftover
# copy cannot quietly look like the live list. See read_sac_students.
SAC_STUDENTS_RETIRED_CSV = SHARED_DIR / "sac-students.csv"
# ---------------------------------------------------------------------------

PRINT_JOB_CSV = (DATA_DIR or SCRIPT_DIR) / "print-job.csv"

# Settings that must be present in print-job.csv, and must be whole numbers.
# Policy rather than form answers, which is why they carry no label there.
# Each is the DEFAULT offered at one of test-admin's questions, not a fixed
# figure: spare_per_section and sac_cushion can be overridden at the prompt, and
# sac_fallback stands in when sac-students.xlsx cannot be read.
REQUIRED_NUMBERS = ("spare_per_section", "sac_fallback", "sac_cushion")


#%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%
#% THE REPORT
#%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%
# NOTHING IS PRINTED WHEN IT HAPPENS. Every line the script has to say is held
# here and emitted in a block, at a moment someone is reading: just before each
# question (the lines so far are what the question is about), and once more at
# the very end. A warning raised while a table is half-built therefore still
# lands under that table, not in the middle of it.
#
# Two lists, not one, because a leading table has to come first however late a
# warning was raised: warnings are often discovered while building the very data
# the table reports on.
report_table = []
report_warnings = []
warning_count = 0

# Colour only for a terminal. Redirected to a file or through a pager, escape
# codes are noise, so both a tty and the absence of NO_COLOR are required.
# NO_COLOR is the informal cross-tool convention: any value at all means off.
USE_COLOUR = sys.stderr.isatty() and "NO_COLOR" not in os.environ

COLOURS = {"red": "\033[31m", "yellow": "\033[33m", "green": "\033[32m",
           "bold": "\033[1m", "dim": "\033[2m"}


def paint(text, colour):
    """Wrap text in an ANSI colour, or return it untouched when not a terminal."""
    if not USE_COLOUR:
        return text
    return f"{COLOURS[colour]}{text}\033[0m"


def say(line):
    """Add a line to the report's leading table."""
    report_table.append(line)


def warn(message):
    """Note a problem without stopping. Nothing here belongs on the sheet."""
    global warning_count
    warning_count += 1
    report_warnings.append(f"  {paint('!', 'yellow')} {message}")


def count_warning():
    """Count a problem whose text is already in the table.

    warn() does two things -- records a message and bumps the tally. This is the
    other half, for a row that belongs in the leading table rather than in the
    warning list. A stale-class-list row is the case that needs it: routing a bad
    row through warn() would put an "!" marker on some rows of a table and not
    others, but it still has to reach the final count or the tally would
    understate the trouble.

    A function rather than `global warning_count' at the call site, because the
    call site is now in another module and the global statement there would
    quietly bind a second, unrelated counter.
    """
    global warning_count
    warning_count += 1


def flush_report(tally=True):
    """Print everything collected since the last flush.

    `tally' adds the running count of warnings as a last line. That line is for
    the END of a run -- it is the verdict -- so the flushes before a question
    pass tally=False and show only the lines themselves. The count is never
    reset: it is the whole run's, whichever flush reports it.

    Also called from die(), so a fatal error still shows the table and any
    warnings gathered before it -- those are usually the context that explains
    the error.
    """
    # stdout is block-buffered when redirected, so lines printed there would
    # otherwise surface after this block rather than before it.
    sys.stdout.flush()
    if report_table or report_warnings:
        print("", file=sys.stderr)
        for line in report_table + report_warnings:
            print(line, file=sys.stderr)
    if tally and warning_count:
        print(
            paint(f"  {warning_count} warning(s) -- check before printing.", "yellow"),
            file=sys.stderr,
        )
    sys.stderr.flush()
    report_table.clear()
    report_warnings.clear()


def die(message):
    flush_report()
    print(paint(f"{PROGRAM}: error: {message}", "red"), file=sys.stderr)
    sys.exit(1)


def ensure_output_dir():
    """OUTPUT_DIR, created if need be -- but only its last level.

    Tests/ must already be there. If it is not, OneDrive is not signed in or not
    synced, and creating the whole chain would build a plausible-looking Team
    folder on the local disk that never reaches anyone. Better to stop and say so.
    """
    if not OUTPUT_DIR.parent.is_dir():
        die(
            f"{OUTPUT_DIR.parent} not found, so there is nowhere to put the "
            f"output. Is OneDrive signed in and synced?"
        )
    OUTPUT_DIR.mkdir(exist_ok=True)
    return OUTPUT_DIR


#%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%
#% THE COURSE'S SHAPE
#%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%

def parse_test_number(value):
    """Check N before it reaches a heading and a filename.

    N is interpolated into test-{N}-attendance-sheet.tex and test-{N}.xlsx,
    so an unchecked value fails late and obscurely: a path lands a directory
    separator in the middle of the filename and the script dies in write_text on
    a directory that was never going to exist. Checking here turns that into one
    line from argparse.

    The mistake worth naming is passing a data file instead -- a reasonable guess
    at "run it on my roster", but every data file these scripts read is found for
    them, and N really is only the number printed on the paperwork.
    """
    if value.isdigit() and int(value) > 0:
        return value

    looks_like_a_file = (
        "/" in value or value.lower().endswith((".csv", ".tsv", ".pdf"))
        or "section-rosters" in value
    )
    if looks_like_a_file:
        raise argparse.ArgumentTypeError(
            f"{value!r} looks like a data file, not a test number. The data "
            f"files are found automatically; N is just the number printed on "
            f"the paperwork, e.g. `{PROGRAM} 2'."
        )
    raise argparse.ArgumentTypeError(
        f"{value!r} is not a positive whole number, e.g. 2"
    )


def read_sections():
    """course-sections.csv -> {section: {...}}, in the order it lists them.

    That order is the page order of the attendance sheets and the booklet order
    of the print job, which is why this keeps an OrderedDict rather than sorting:
    the CSV is already FR01A..FR24A and it is the file a human would edit if that
    ever needed to change.

    THE DAY IS CARRIED RAW, not normalised or validated here. This function's job
    is to report what the CSV says. test-admin's split_by_day, which divides the
    whole order on it, is the right place to be strict, because that is where a
    misspelt day would silently under-order.
    """
    if DATA_DIR is None:
        die(
            f"course-sections.csv not found in {DATA_DIR_NAME}/ at the top "
            f"of the repository ({REPO_ROOT}). It is what says which sections "
            "exist, so there is nothing to build without it."
        )
    if not SECTIONS_CSV.exists():
        die(f"{SECTIONS_CSV} not found")

    sections = OrderedDict()
    with SECTIONS_CSV.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            code = (row.get("Section") or "").strip()
            if not code:
                continue
            # The Enrolment column is "22/25" -- registered over capacity. Only
            # the registered half is of interest: it is what the order is
            # computed from, and what the class lists are checked against.
            enrolment = (row.get("Enrolment") or "").strip()
            registered = None
            if "/" in enrolment:
                head = enrolment.split("/", 1)[0].strip()
                if head.isdigit():
                    registered = int(head)
            elif enrolment.isdigit():
                registered = int(enrolment)
            sections[code] = {
                "ta": (row.get("TA") or "").strip(),
                "instructor": (row.get("Instructor") or "").strip(),
                "day": (row.get("Day") or "").strip(),
                "room": (row.get("Room") or "").strip(),
                "registered": registered,
            }

    if not sections:
        die(f"{SECTIONS_CSV} has no section rows")
    return sections


#%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%
#% THE PRINT JOB'S FIXED ANSWERS
#%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%

def read_print_job():
    """print-job.csv -> (form fields in file order, settings by name).

    THE FILE IS THE FORM. Its rows are in the order the web form asks its
    questions, and each carries the question's own wording in `label', so when
    Print Services rearranges or renames a field the fix is one line of CSV and
    no code. That is also why the order here is the order of the output: reading
    down the generated file and down the web page should be the same motion.

    A row with an EMPTY LABEL is not a question -- it is a setting (a policy
    number, a URL) that the script or the reader wants but the form never asks
    for. That is the whole distinction; nothing else separates the two kinds.
    """
    if not PRINT_JOB_CSV.exists():
        die(
            f"{PRINT_JOB_CSV} not found. It holds the form's fixed answers -- "
            f"your name, the account number, the pickup choice -- and the "
            f"policy numbers ({', '.join(REQUIRED_NUMBERS)})."
        )

    fields, settings = [], {}
    with PRINT_JOB_CSV.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        missing = {"field", "label", "value"} - set(reader.fieldnames or [])
        if missing:
            die(
                f"{PRINT_JOB_CSV.name} is missing column(s) "
                f"{', '.join(sorted(missing))}; it needs field, label, value."
            )
        for row in reader:
            name = (row.get("field") or "").strip()
            if not name:
                continue
            label = (row.get("label") or "").strip()
            value = (row.get("value") or "").strip()
            settings[name] = value
            if label:
                fields.append((name, label, value))

    for name in REQUIRED_NUMBERS:
        if name not in settings:
            die(f"{PRINT_JOB_CSV.name} has no `{name}' row")
        if not settings[name].isdigit():
            die(
                f"{PRINT_JOB_CSV.name}: `{name}' is {settings[name]!r}, which is "
                f"not a whole number"
            )
    return fields, settings


#%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%
#% WHO WRITES AT THE SAC
#%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%%
# REGISTERING WITH THE SAC IS NOT WRITING THERE. Some students who register keep
# writing tests in their own tutorial room and use the SAC only for the final
# exam. That is the whole reason sac-students.xlsx has a `writes-at-sac' column
# rather than being a bare list of names: presence in the file says nothing about
# where a booklet has to be.
#
# The column is filled in from experience -- nobody knows before a student's
# first test -- and once a student writes at the SAC they carry on doing so, so
# the file converges rather than needing a fresh census every test.
#
# THREE STATES, and the middle one is the point:
#
#   yes      counted at the SAC and NOT in their section. The saving.
#   no       no effect whatever; they sit in their section like anyone else.
#   unknown  counted in BOTH -- the section keeps its booklet AND the SAC gets
#            one. Deliberately over-printing by one sheet per unknown, because
#            the alternative is a student arriving at the SAC to no paper. It is
#            also exactly what the flat ballpark used to buy, now spent only on
#            the people it is genuinely uncertain for.

SAC_YES = {"yes", "y", "true", "1"}
SAC_NO = {"no", "n", "false", "0"}

# Spelled-out unknowns. An EMPTY cell already means unknown and always has; this
# set is for the ones people type on purpose, and `unknown' in particular is an
# option in the workbook's dropdown (see make-sac-workbook). Without this they
# would fall to the else-branch below and be warned about as unrecognised, which
# would put a warning on every row that is honestly answered "not yet".
SAC_UNKNOWN = {"unknown", "?", "tbd", "not yet"}

SAC_COLUMNS = ("last", "first", "section", "writes-at-sac")


def fold_name(text):
    """Casefold and strip accents, for comparing two spellings of one name.

    Used to de-duplicate rows here and, in test-admin, to match a row
    against the registrar's class list. Comparison only -- nothing folded by this
    is ever printed, so it is free to be lossy.
    """
    decomposed = unicodedata.normalize("NFKD", text)
    bare = "".join(c for c in decomposed if not unicodedata.combining(c))
    return " ".join(bare.casefold().split())


def read_sac_students(sections):
    """sac-students.xlsx -> (counts, marked, available).

        counts[section]  {"yes": n, "unknown": n}    what the arithmetic needs
        marked[section]  [(last, first), ...]        the `yes' rows, to mark on
                                                     the attendance sheet
        available        False when the file could not be read at all

    NOT FATAL WHEN ABSENT -- OneDrive may simply not be there. The caller falls
    back to print-job.csv's sac_fallback, which is the old flat ballpark: it
    over-prints, so a missing file wastes paper rather than stranding anyone.
    That is why this warns instead of dying.

    NOTHING IS DROPPED SILENTLY. A row whose section is not in
    course-sections.csv cannot be attributed to a lecture day and so cannot be
    ordered for; it is reported by name-less description and left out, because
    the fix is one character in a shared file and the alternative is a number
    nobody can account for.

    A WORKBOOK RATHER THAN A CSV because three other people edit it in Teams;
    see the PORTING block above. Only the READING changed in that move -- every
    rule below it is the one the CSV had, and the row numbers quoted in warnings
    are now real spreadsheet row numbers, so `Ctrl-G 14' lands on the bad row.
    """
    # A leftover CSV is NOT read, only reported. Two files that could each look
    # like the live list is the same trap as a stale local copy standing in for
    # the shared one, and it is worth a line of noise to close it.
    if SAC_STUDENTS_RETIRED_CSV.exists():
        warn(
            f"{SAC_STUDENTS_RETIRED_CSV.name} is still in the shared folder. "
            f"It is NOT read -- {SAC_STUDENTS_XLSX.name} replaced it. Delete "
            f"the .csv so nobody edits the wrong file."
        )

    if not SAC_STUDENTS_XLSX.exists():
        warn(
            f"{SAC_STUDENTS_XLSX.name} not found in the shared folder, so the "
            f"SAC order falls back to the flat ballpark and no section is "
            f"reduced. (Is OneDrive signed in?)"
        )
        return {}, {}, False

    # Only reachable when the script was started some way that skipped uv, since
    # the inline dependency block at the top of each one asks for openpyxl. Fatal
    # rather than a fallback: the file is RIGHT THERE and readable, so quietly
    # over-printing instead of reading it would be a lie about why.
    if load_workbook is None:
        die(
            f"openpyxl is not installed, so {SAC_STUDENTS_XLSX.name} cannot be "
            f"read. Run this script directly (./{PROGRAM}) and uv will fetch "
            f"it, or use `uv run --script {PROGRAM}'."
        )

    # data_only so a cell holding a formula gives its cached value rather than
    # the formula text; read_only because nothing here writes and the sheet may
    # be open in Excel at the other end of OneDrive. A file caught mid-sync
    # raises rather than returning nonsense, and that is the case this catches:
    # treat it exactly like a missing file, which over-prints.
    try:
        workbook = load_workbook(
            SAC_STUDENTS_XLSX, read_only=True, data_only=True
        )
    except Exception as problem:
        warn(
            f"{SAC_STUDENTS_XLSX.name} could not be opened ({problem}). It may "
            f"be part-way through syncing, or open for editing. Falling back "
            f"to the ballpark, which over-prints."
        )
        return {}, {}, False

    counts, marked, seen = {}, {}, set()
    try:
        rows = workbook.active.iter_rows(min_row=1, values_only=True)

        # BY HEADER NAME, NEVER BY POSITION -- the same rule DictReader gave the
        # CSV, and it matters more now: four people share this sheet and someone
        # will eventually insert a column. Folded so a stray capital or a
        # trailing space typed into a header does not fail the whole file.
        header = next(rows, None) or ()
        column = {
            str(name).strip().casefold(): index
            for index, name in enumerate(header)
            if name is not None
        }
        missing = set(SAC_COLUMNS) - set(column)
        if missing:
            warn(
                f"{SAC_STUDENTS_XLSX.name} is missing column(s) "
                f"{', '.join(sorted(missing))}; it needs "
                f"{', '.join(SAC_COLUMNS)}. Falling back to the ballpark."
            )
            return {}, {}, False

        def field(row, name):
            """One cell as clean text: blank for empty, str() for a number."""
            index = column[name]
            value = row[index] if index < len(row) else None
            return "" if value is None else str(value).strip()

        # start=2 because the header was row 1, so `row' counts as Excel counts.
        for number, row in enumerate(rows, start=2):
            last = field(row, "last")
            first = field(row, "first")
            section = field(row, "section")
            # Also throws away the run of empty rows Excel leaves below the data:
            # a saved sheet routinely reports more rows than anyone typed into.
            if not (last or first or section):
                continue

            if section not in sections:
                warn(
                    f"{SAC_STUDENTS_XLSX.name} row {number}: section "
                    f"{section!r} is not in course-sections.csv, so this "
                    f"student is not counted in any day's order. Fix the "
                    f"section code."
                )
                continue

            key = (fold_name(last), fold_name(first), section)
            if key in seen:
                warn(
                    f"{SAC_STUDENTS_XLSX.name} row {number}: a second row for "
                    f"the same student in {section}; counted once."
                )
                continue
            seen.add(key)

            state = field(row, "writes-at-sac").casefold()
            if state in SAC_YES:
                where = "yes"
            elif state in SAC_NO:
                where = "no"
            elif state == "" or state in SAC_UNKNOWN:
                where = "unknown"
            else:
                warn(
                    f"{SAC_STUDENTS_XLSX.name} row {number}: writes-at-sac is "
                    f"{state!r}, which is neither yes nor no. Treating it as "
                    f"not yet known, which orders a booklet in both places."
                )
                where = "unknown"

            tally = counts.setdefault(section, {"yes": 0, "unknown": 0})
            if where == "yes":
                tally["yes"] += 1
                marked.setdefault(section, []).append((last, first))
            elif where == "unknown":
                tally["unknown"] += 1
    finally:
        # A read_only workbook holds the file open until told otherwise.
        workbook.close()

    return counts, marked, True
